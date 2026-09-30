import hashlib
import json
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from reviews.models import Course, CurriculumVersion, Program
from reviews.management.import_utils import resolve_columns


def _text(value):
    return "" if value is None else str(value).strip()


class Command(BaseCommand):
    help = "Import one explicitly mapped curriculum worksheet. Source workbook remains read-only."

    def add_arguments(self, parser):
        parser.add_argument("--workbook", required=True)
        parser.add_argument("--config", required=True, help="JSON: program_code/name, sheet, version_label, header_row, headers")
        parser.add_argument("--dry-run", action="store_true")

    def handle(self, *args, **options):
        path = Path(options["workbook"]).expanduser().resolve()
        if not path.is_file():
            raise CommandError("Workbook path does not exist.")
        try:
            config = json.loads(Path(options["config"]).read_text(encoding="utf-8"))
        except Exception as exc:
            raise CommandError("Could not read JSON import config.") from exc
        required = {"program_code", "program_name", "sheet", "version_label", "header_row", "headers"}
        if not required.issubset(config) or not {"course_code", "course_name"}.issubset(config["headers"]):
            raise CommandError("Config must define program, sheet, version, header row, course_code and course_name.")
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        workbook = load_workbook(path, read_only=True, data_only=False)
        if config["sheet"] not in workbook.sheetnames:
            workbook.close()
            raise CommandError("Configured worksheet was not found.")
        sheet = workbook[config["sheet"]]
        header_row = int(config["header_row"])
        columns, missing, by_header = resolve_columns(sheet, header_row, config["headers"])
        if missing:
            workbook.close()
            raise CommandError(f"Configured source headers not found: {', '.join(missing)}")
        rows = []
        for row_index, values in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            mapped = {field: values[columns[field] - 1] if len(values) >= columns[field] else None
                      for field in config["headers"]}
            if not _text(mapped.get("course_code")) and not _text(mapped.get("course_name")):
                continue
            rows.append((row_index, mapped))
        workbook.close()
        if options["dry_run"]:
            self.stdout.write(f"DRY RUN: {len(rows)} course rows; source sha256={digest}")
            return
        with transaction.atomic():
            program, _ = Program.objects.get_or_create(code=config["program_code"], defaults={"name": config["program_name"]})
            if program.name != config["program_name"]:
                raise CommandError("Program code exists with a different name. Resolve manually before import.")
            curriculum, created = CurriculumVersion.objects.get_or_create(
                source_sha256=digest, source_sheet=config["sheet"],
                defaults={"program": program, "version_label": config["version_label"],
                          "source_file": path.name, "source_sheet": config["sheet"]},
            )
            if not created:
                self.stdout.write(f"Already imported: {curriculum.pk}")
                return
            for row_index, mapped in rows:
                raw_code, raw_name = _text(mapped.get("course_code")), _text(mapped.get("course_name"))
                credit_value = mapped.get("credits")
                try:
                    credits = Decimal(str(credit_value)) if isinstance(credit_value, (int, float, Decimal)) else None
                except InvalidOperation:
                    credits = None
                source_cells = {}
                for field, source_header in config["headers"].items():
                    col = columns[field]
                    source_cells[field] = {"cell": f"{get_column_letter(col)}{row_index}", "header": _text(sheet.cell(header_row, col).value),
                                            "raw_value": _text(mapped.get(field))}
                Course.objects.create(
                    curriculum=curriculum, raw_code=raw_code, raw_name=raw_name, credits=credits,
                    course_type=_text(mapped.get("course_type")), assessment=_text(mapped.get("assessment")),
                    source_row=row_index, source_cells=source_cells,
                )
        self.stdout.write(f"Imported {len(rows)} rows as DRAFT. Curriculum id={curriculum.pk}")
