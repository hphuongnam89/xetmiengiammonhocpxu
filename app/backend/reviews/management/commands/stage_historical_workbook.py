import hashlib
import json
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from reviews.management.import_utils import resolve_columns
from reviews.models import AuditEvent, HistoricalDecision


def _text(value):
    return "" if value is None else str(value).strip()


def _number(value):
    if value in (None, ""):
        return Decimal("0"), True
    try:
        return Decimal(str(value).replace(",", ".").strip()), True
    except (InvalidOperation, ValueError):
        return Decimal("0"), False


class Command(BaseCommand):
    help = "Stage explicitly configured worksheet decisions as pending; anonymizes worksheet names."

    def add_arguments(self, parser):
        parser.add_argument("--workbook", required=True)
        parser.add_argument("--config", required=True)
        parser.add_argument("--apply", action="store_true", help="Persist pending rows; default is preview only.")

    def handle(self, *args, **options):
        path = Path(options["workbook"]).expanduser().resolve()
        if not path.is_file():
            raise CommandError("Workbook path does not exist.")
        try:
            config = json.loads(Path(options["config"]).read_text(encoding="utf-8"))
        except Exception as exc:
            raise CommandError("Could not read JSON import config.") from exc
        if not {"header_row", "data_start_row", "program_code", "headers"}.issubset(config):
            raise CommandError("Config must define header_row, data_start_row, program_code and headers.")
        if not {"course_code", "course_name", "full_credit", "partial_credit"}.issubset(config["headers"]):
            raise CommandError("Config must map course_code, course_name, full_credit and partial_credit.")

        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        workbook = load_workbook(path, read_only=True, data_only=True)
        staged = []
        matched_sheets = skipped_sheets = ambiguous_rows = 0
        full_rows = partial_rows = 0
        header_row = int(config["header_row"])
        start_row = int(config["data_start_row"])
        for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
            try:
                columns, missing, _ = resolve_columns(sheet, header_row, config["headers"])
            except (StopIteration, ValueError):
                missing = ["unreadable header"]
            if missing:
                skipped_sheets += 1
                continue
            matched_sheets += 1
            safe_sheet_id = f"worksheet-{sheet_index:03d}"
            legacy_sheet_id = f"worksheet-{sheet_index:03d}-{hashlib.sha256(sheet.title.encode()).hexdigest()[:8]}"
            for row_number, values in enumerate(sheet.iter_rows(min_row=start_row, values_only=True), start=start_row):
                def cell(field):
                    index = columns[field] - 1
                    return values[index] if index < len(values) else None

                course_code, course_name = _text(cell("course_code")), _text(cell("course_name"))
                full_raw, partial_raw = cell("full_credit"), cell("partial_credit")
                full_value, full_valid = _number(full_raw)
                partial_value, partial_valid = _number(partial_raw)
                if not course_code or not course_name or (full_raw in (None, "", 0) and partial_raw in (None, "", 0)):
                    continue
                full_positive, partial_positive = full_value > 0, partial_value > 0
                decision = "FULL" if full_positive and not partial_positive else "PARTIAL" if partial_positive and not full_positive else ""
                if decision == "FULL":
                    full_rows += 1
                elif decision == "PARTIAL":
                    partial_rows += 1
                else:
                    ambiguous_rows += 1
                cells = {}
                for field in ("course_code", "course_name", "full_credit", "partial_credit"):
                    col = columns[field]
                    cells[field] = {
                        "cell": f"{get_column_letter(col)}{row_number}",
                        "header": _text(sheet.cell(header_row, col).value),
                        "raw_value": _text(cell(field)),
                    }
                raw_decision = json.dumps(
                    {"full_credit": _text(full_raw), "partial_credit": _text(partial_raw),
                     "numeric_values_valid": full_valid and partial_valid}, ensure_ascii=False,
                )
                staged.append({
                    "source_sheet": safe_sheet_id, "source_row": row_number,
                    "legacy_sheet_id": legacy_sheet_id,
                    "course_code": course_code[:100], "course_name": course_name[:500],
                    "decision": decision, "decision_raw": raw_decision[:255], "source_cells": cells,
                })
        workbook.close()
        self.stdout.write(
            f"Preview: {matched_sheets} mapped sheets, {skipped_sheets} skipped; "
            f"{len(staged)} rows staged ({full_rows} FULL, {partial_rows} PARTIAL, {ambiguous_rows} needs manual mapping)."
        )
        if not options["apply"]:
            self.stdout.write("No database changes. Re-run with --apply after checking the mapping summary.")
            return
        with transaction.atomic():
            created = anonymized = 0
            for row in staged:
                legacy = HistoricalDecision.objects.filter(
                    source_sha256=digest, source_sheet=row["legacy_sheet_id"], source_row=row["source_row"],
                    status="PENDING",
                ).first()
                if legacy:
                    legacy.source_sheet = row["source_sheet"]
                    legacy.save(update_fields=["source_sheet"])
                    anonymized += 1
                _, was_created = HistoricalDecision.objects.get_or_create(
                    source_sha256=digest, source_sheet=row["source_sheet"], source_row=row["source_row"],
                    defaults={
                        "source_file": path.name,
                        "program_code": str(config["program_code"])[:100],
                        "course_code": row["course_code"], "course_name": row["course_name"],
                        "decision_raw": row["decision_raw"], "decision": row["decision"],
                        "source_cells": row["source_cells"],
                    },
                )
                created += int(was_created)
            if created or anonymized:
                AuditEvent.objects.create(
                    actor=None, action="HISTORICAL_WORKBOOK_STAGED_OR_ANONYMIZED", entity_type="HistoricalWorkbook",
                    entity_id=digest[:32],
                    payload={"source_file": path.name, "source_sha256": digest,
                             "mapped_sheets": matched_sheets, "skipped_sheets": skipped_sheets,
                             "new_pending_rows": created, "full_rows": full_rows,
                             "partial_rows": partial_rows, "ambiguous_rows": ambiguous_rows,
                             "sheet_labels_anonymized": anonymized, "student_sheet_titles_stored": False},
                )
        self.stdout.write(f"Persisted {created} new records as PENDING. No precedent was approved.")
