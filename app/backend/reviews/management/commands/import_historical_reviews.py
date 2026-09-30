import hashlib
import json
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from reviews.models import HistoricalDecision
from reviews.management.import_utils import resolve_columns


def _text(value):
    return "" if value is None else str(value).strip()


class Command(BaseCommand):
    help = "Stage explicitly mapped historical decisions as pending review; never auto-approve imported rows."

    def add_arguments(self, parser):
        parser.add_argument("--workbook", required=True)
        parser.add_argument("--config", required=True)
        parser.add_argument("--dry-run", action="store_true")

    def handle(self, *args, **options):
        path = Path(options["workbook"]).expanduser().resolve()
        if not path.is_file():
            raise CommandError("Workbook path does not exist.")
        try:
            config = json.loads(Path(options["config"]).read_text(encoding="utf-8"))
        except Exception as exc:
            raise CommandError("Could not read JSON import config.") from exc
        if not {"sheet", "header_row", "headers"}.issubset(config):
            raise CommandError("Config must define sheet, header_row and headers.")
        if "course_code" not in config["headers"] or ("decision" not in config["headers"] and not config.get("decision_columns")):
            raise CommandError("Config must map course_code and decision or decision_columns.")
        allowed_decisions = {"FULL", "PARTIAL", "NOT_ELIGIBLE"}
        if not set(str(value).strip().upper() for value in config.get("decision_map", {}).values()).issubset(allowed_decisions):
            raise CommandError("decision_map values must be FULL, PARTIAL or NOT_ELIGIBLE.")
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        workbook = load_workbook(path, read_only=True, data_only=False)
        if config["sheet"] not in workbook.sheetnames:
            workbook.close()
            raise CommandError("Configured worksheet was not found.")
        sheet = workbook[config["sheet"]]
        header_row = int(config["header_row"])
        field_specs = dict(config["headers"])
        for field, spec in config.get("decision_columns", {}).items():
            field_specs[f"outcome_{field}"] = spec
        header_columns, missing, by_header = resolve_columns(sheet, header_row, field_specs)
        if missing:
            workbook.close()
            raise CommandError(f"Configured source headers not found: {', '.join(missing)}")
        records = []
        for row_number, values in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            mapped = {field: values[header_columns[field] - 1] if len(values) >= header_columns[field] else None
                      for field in config["headers"]}
            code, raw_decision = _text(mapped.get("course_code")), _text(mapped.get("decision"))
            if not code:
                continue
            outcomes_exist = any(
                _text(values[header_columns[f"outcome_{outcome}"] - 1] if len(values) >= header_columns[f"outcome_{outcome}"] else None)
                for outcome in config.get("decision_columns", {})
            )
            if config.get("decision_columns") and not outcomes_exist:
                continue
            if not config.get("decision_columns") and not raw_decision:
                continue
            outcome_values = {}
            for outcome, _spec in config.get("decision_columns", {}).items():
                value = values[header_columns[f"outcome_{outcome}"] - 1] if len(values) >= header_columns[f"outcome_{outcome}"] else None
                outcome_values[outcome] = _text(value)
            if "decision" in config["headers"]:
                decision_value = raw_decision
            else:
                decision_value = json.dumps(outcome_values, ensure_ascii=False)
            normalized_map = {str(key).strip().casefold(): str(value).strip().upper() for key, value in config.get("decision_map", {}).items()}
            normalized = normalized_map.get(decision_value.casefold(), "")
            cells = {}
            for field, spec in field_specs.items():
                if field == "student_ref":
                    continue
                value = values[header_columns[field] - 1] if len(values) >= header_columns[field] else None
                col = header_columns[field]
                header_text = _text(sheet.cell(header_row, col).value)
                source_field = field.removeprefix("outcome_") if field.startswith("outcome_") else field
                cells[source_field] = {"cell": f"{get_column_letter(col)}{row_number}", "header": header_text, "raw_value": _text(value)}
            records.append({"row": row_number, "mapped": mapped, "raw_decision": raw_decision,
                            "decision": normalized, "cells": cells, "decision_raw": decision_value})
        workbook.close()
        if options["dry_run"]:
            known = sum(bool(record["decision"]) for record in records)
            self.stdout.write(f"DRY RUN: {len(records)} rows, {known} mapped decisions, all require review; sha256={digest}")
            return
        with transaction.atomic():
            created = 0
            for record in records:
                mapped = record["mapped"]
                student_ref = _text(mapped.get("student_ref"))
                student_hash = hashlib.sha256(student_ref.encode()).hexdigest() if student_ref else ""
                _, was_created = HistoricalDecision.objects.get_or_create(
                    source_sha256=digest, source_sheet=config["sheet"], source_row=record["row"],
                    defaults={
                        "source_file": path.name, "source_key_hash": student_hash,
                        "program_code": _text(mapped.get("program_code", config.get("program_code", ""))),
                        "course_code": _text(mapped.get("course_code"))[:100],
                        "course_name": _text(mapped.get("course_name"))[:500],
                        "decision_raw": record["decision_raw"][:255], "decision": record["decision"],
                        "source_cells": record["cells"],
                    },
                )
                created += int(was_created)
        self.stdout.write(f"Staged {created} rows as PENDING. No historical decision was auto-approved.")
